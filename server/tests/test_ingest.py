import io

import pandas as pd

from execflow.ingest import ingest, load_dataset, save_dataset
from execflow.profile import candidate_metrics, profile_dataset


def test_excel_header_detection_and_coercion(dataset):
    t = dataset.table("KPI_Data")
    assert t.columns == ["Month", "System", "Failures", "Delay Minutes", "PM Completion"]
    assert t.row_count == 6
    assert pd.api.types.is_datetime64_any_dtype(t.df["Month"])
    assert pd.api.types.is_numeric_dtype(t.df["Delay Minutes"])  # "1,050" became a number
    assert t.df["Delay Minutes"].sum() == 2380
    assert pd.api.types.is_numeric_dtype(t.df["PM Completion"])  # "98%" became 98


def test_small_sheet_is_kept(dataset):
    assert {t.name for t in dataset.tables} == {"KPI_Data", "Cover"}


def test_csv(csv_bytes):
    ds = ingest("contracts.csv", csv_bytes)
    t = ds.table("data")
    assert t.row_count == 3
    assert pd.api.types.is_numeric_dtype(t.df["Value AED"])
    assert pd.api.types.is_datetime64_any_dtype(t.df["Due"])


def test_docx_tables_and_text():
    import docx

    d = docx.Document()
    d.add_paragraph("Monthly report narrative. Service affecting failures increased in June.")
    tbl = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("KPI", "Value"), ("TSR", "2.1"), ("TSA", "99.2")]):
        tbl.rows[r].cells[0].text, tbl.rows[r].cells[1].text = a, b
    buf = io.BytesIO()
    d.save(buf)
    ds = ingest("report.docx", buf.getvalue())
    assert ds.text and "failures increased" in ds.text[0]
    assert ds.table("table1").columns == ["KPI", "Value"]
    assert float(ds.table("table1").df["Value"].sum()) == 101.3


def test_unsupported_and_empty():
    import pytest

    with pytest.raises(ValueError):
        ingest("x.exe", b"")


def test_roundtrip(dataset, tmp_root):
    save_dataset(dataset, tmp_root)
    back = load_dataset(dataset.id, tmp_root)
    assert back.table("KPI_Data").row_count == 6
    assert pd.api.types.is_datetime64_any_dtype(back.table("KPI_Data").df["Month"])


def test_profile(dataset):
    p = profile_dataset(dataset)
    kpi = next(t for t in p["tables"] if t["name"] == "KPI_Data")
    kinds = {c["name"]: c["kind"] for c in kpi["columns"]}
    assert kinds["Month"] == "date" and kinds["System"] == "category" and kinds["Failures"] == "numeric"
    system = next(c for c in kpi["columns"] if c["name"] == "System")
    assert {t["value"] for t in system["top"]} == {"Signalling", "Rolling Stock"}
    cands = candidate_metrics(p)
    assert cands[0]["table"] == "KPI_Data" and "Failures" in cands[0]["metrics"] and cands[0]["date_columns"] == ["Month"]
