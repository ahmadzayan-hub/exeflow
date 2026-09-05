"""File ingestion: xlsx / csv / pdf / docx into named tables plus text.

Every number the agents later quote must be recomputable from these tables,
so parsing is deterministic and nothing is inferred by a model here.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

SUPPORTED = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".pdf", ".docx"}
MAX_TEXT_CHUNK = 1500


@dataclass
class Table:
    name: str
    df: pd.DataFrame

    @property
    def row_count(self) -> int:
        return int(len(self.df))

    @property
    def columns(self) -> list[str]:
        return [str(c) for c in self.df.columns]


@dataclass
class Dataset:
    id: str
    name: str
    filename: str
    kind: str
    tables: list[Table] = field(default_factory=list)
    text: list[str] = field(default_factory=list)
    created_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat(timespec="seconds"))

    def table(self, name: str) -> Table:
        for t in self.tables:
            if t.name == name:
                return t
        raise KeyError(f"unknown table '{name}'; available: {[t.name for t in self.tables]}")

    def meta(self) -> dict:
        return {
            "id": self.id,
            "name": self.name,
            "filename": self.filename,
            "kind": self.kind,
            "created_at": self.created_at,
            "tables": [{"name": t.name, "columns": t.columns, "row_count": t.row_count} for t in self.tables],
            "text_chunks": len(self.text),
        }


def slug(s: str) -> str:
    s = re.sub(r"[^A-Za-z0-9؀-ۿ]+", "_", str(s)).strip("_")
    return s[:60] or "table"


def _clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    cols = []
    seen: dict[str, int] = {}
    for c in df.columns:
        name = str(c).strip()
        if not name or name.lower().startswith("unnamed"):
            name = "col"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        cols.append(name)
    df = df.copy()
    df.columns = cols
    return df.dropna(axis=1, how="all").dropna(axis=0, how="all")


def _find_header(raw: pd.DataFrame) -> int:
    """Row index that looks like a header: mostly non-null strings, in the first 15 rows."""
    best, best_score = 0, -1.0
    for i in range(min(15, len(raw))):
        row = raw.iloc[i]
        non_null = row.notna().sum()
        if non_null == 0:
            continue
        strings = sum(isinstance(v, str) and v.strip() != "" for v in row)
        score = strings / max(1, len(row)) + (non_null / max(1, len(row))) * 0.5
        if score > best_score:
            best, best_score = i, score
    return best


def _frame_from_raw(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return raw
    h = _find_header(raw)
    header = [str(v) if pd.notna(v) else "col" for v in raw.iloc[h]]
    df = raw.iloc[h + 1 :].reset_index(drop=True)
    df.columns = header
    df = _clean_columns(df)
    return df.convert_dtypes().infer_objects()


def _coerce(df: pd.DataFrame) -> pd.DataFrame:
    """Numbers stored as text ("1,250", "12%") and ISO-like dates become typed columns."""
    out = df.copy()
    for c in out.columns:
        s = out[c]
        if s.dtype == object or pd.api.types.is_string_dtype(s):
            txt = s.astype("string").str.strip()
            num = pd.to_numeric(txt.str.replace(",", "", regex=False).str.rstrip("%"), errors="coerce")
            present = int(txt.notna().sum())
            if present and num.notna().sum() >= max(1, 0.8 * present):
                out[c] = num
                continue
            dt = pd.to_datetime(txt, errors="coerce", format="mixed")
            if present and dt.notna().sum() >= max(1, 0.8 * present):
                out[c] = dt
    return out


def read_excel(data: bytes) -> list[Table]:
    book = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None)
    tables = []
    for sheet, raw in book.items():
        df = _frame_from_raw(raw)
        if df.empty or len(df.columns) < 1:
            continue
        tables.append(Table(slug(sheet), _coerce(df)))
    return tables


def read_csv(data: bytes, sep: str | None = None) -> list[Table]:
    text = data.decode("utf-8-sig", errors="replace")
    raw = pd.read_csv(io.StringIO(text), sep=sep, engine="python", header=None)
    return [Table("data", _coerce(_frame_from_raw(raw)))]


def _chunks(paragraphs: list[str]) -> list[str]:
    out, buf = [], ""
    for p in paragraphs:
        p = p.strip()
        if not p:
            continue
        if len(buf) + len(p) > MAX_TEXT_CHUNK and buf:
            out.append(buf)
            buf = p
        else:
            buf = f"{buf}\n{p}" if buf else p
    if buf:
        out.append(buf)
    return out


def read_pdf(data: bytes) -> tuple[list[Table], list[str]]:
    import pdfplumber

    tables: list[Table] = []
    paragraphs: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            txt = page.extract_text() or ""
            paragraphs.append(f"[page {i}] {txt}")
            for j, rows in enumerate(page.extract_tables() or [], start=1):
                if not rows or len(rows) < 2:
                    continue
                raw = pd.DataFrame(rows)
                df = _coerce(_frame_from_raw(raw))
                if not df.empty:
                    tables.append(Table(f"page{i}_table{j}", df))
    return tables, _chunks(paragraphs)


def read_docx(data: bytes) -> tuple[list[Table], list[str]]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs]
    tables: list[Table] = []
    for i, t in enumerate(doc.tables, start=1):
        rows = [[c.text for c in r.cells] for r in t.rows]
        if len(rows) < 2:
            continue
        df = _coerce(_frame_from_raw(pd.DataFrame(rows)))
        if not df.empty:
            tables.append(Table(f"table{i}", df))
    return tables, _chunks(paragraphs)


def ingest(filename: str, data: bytes, name: str | None = None) -> Dataset:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type '{ext}'; supported: {sorted(SUPPORTED)}")
    ds = Dataset(id=uuid.uuid4().hex[:12], name=name or Path(filename).stem, filename=filename, kind=ext.lstrip("."))
    if ext in {".xlsx", ".xlsm", ".xls"}:
        ds.tables = read_excel(data)
    elif ext == ".csv":
        ds.tables = read_csv(data, sep=",")
    elif ext == ".tsv":
        ds.tables = read_csv(data, sep="\t")
    elif ext == ".pdf":
        ds.tables, ds.text = read_pdf(data)
    elif ext == ".docx":
        ds.tables, ds.text = read_docx(data)
    if not ds.tables and not ds.text:
        raise ValueError("no tables or text could be extracted from the file")
    return ds


# ---------------------------------------------------------------- persistence

def save_dataset(ds: Dataset, root: Path) -> Path:
    d = root / ds.id
    d.mkdir(parents=True, exist_ok=True)
    (d / "meta.json").write_text(json.dumps(ds.meta(), ensure_ascii=False, indent=2), encoding="utf-8")
    (d / "text.json").write_text(json.dumps(ds.text, ensure_ascii=False), encoding="utf-8")
    for t in ds.tables:
        t.df.to_csv(d / f"{t.name}.csv", index=False)
    return d


def load_dataset(ds_id: str, root: Path) -> Dataset:
    d = root / ds_id
    meta = json.loads((d / "meta.json").read_text(encoding="utf-8"))
    ds = Dataset(id=meta["id"], name=meta["name"], filename=meta["filename"], kind=meta["kind"], created_at=meta["created_at"])
    ds.text = json.loads((d / "text.json").read_text(encoding="utf-8")) if (d / "text.json").exists() else []
    for t in meta["tables"]:
        df = pd.read_csv(d / f"{t['name']}.csv")
        ds.tables.append(Table(t["name"], _coerce(df)))
    return ds


def list_datasets(root: Path) -> list[dict]:
    out = []
    if not root.exists():
        return out
    for d in sorted(root.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        m = d / "meta.json"
        if m.exists():
            out.append(json.loads(m.read_text(encoding="utf-8")))
    return out
